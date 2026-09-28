"""DOCX renderer (python-docx), same structure as the Markdown and HTML reports.

Citations are internal hyperlinks to a bookmark on each finding's heading (``F_0042``; Word
bookmark names cannot contain hyphens). Locations link to the repository at the commit.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.document import Document as DocumentT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from lantern_report.dpia import Report, Section, Table
from lantern_report.narrative import CITATION
from lantern_report.render import DISCLAIMER

LINK_COLOR = "1F5F8B"


def bookmark_name(finding_id: str) -> str:
    return finding_id.replace("-", "_")


class _Writer:
    def __init__(self) -> None:
        self.doc: DocumentT = Document()
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        self._bookmark_id = 0

    # --- inline helpers -----------------------------------------------------------------

    def _link_run(self, text: str) -> Any:
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), LINK_COLOR)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        props.append(color)
        props.append(underline)
        run.append(props)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        return run

    def link(
        self, paragraph: Paragraph, text: str, url: str | None = None, anchor: str | None = None
    ) -> None:
        if not url and not anchor:
            paragraph.add_run(text)
            return
        hyperlink = OxmlElement("w:hyperlink")
        if url:
            rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), rel)
        else:
            hyperlink.set(qn("w:anchor"), anchor or "")
        hyperlink.append(self._link_run(text))
        paragraph._p.append(hyperlink)

    def text_with_citations(self, paragraph: Paragraph, text: str) -> None:
        pos = 0
        for m in CITATION.finditer(text):
            if m.start() > pos:
                paragraph.add_run(text[pos : m.start()])
            self.link(paragraph, f"[{m.group(1)}]", anchor=bookmark_name(m.group(1)))
            pos = m.end()
        if pos < len(text):
            paragraph.add_run(text[pos:])

    def bookmark(self, paragraph: Paragraph, name: str) -> None:
        self._bookmark_id += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bookmark_id))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    # --- blocks ---------------------------------------------------------------------------

    def para(self, text: str = "", italic: bool = False, style: str | None = None) -> Paragraph:
        paragraph = self.doc.add_paragraph(style=style)
        if text:
            self.text_with_citations(paragraph, text)
        if italic:
            for run in paragraph.runs:
                run.italic = True
        return paragraph

    def table(self, table: Table) -> None:
        if table.caption:
            self.para(table.caption, italic=True)
        if not table.rows:
            self.para("Nothing to list.", italic=True)
            return
        grid = self.doc.add_table(rows=1, cols=len(table.columns))
        grid.style = "Table Grid"
        for cell, name in zip(grid.rows[0].cells, table.columns, strict=True):
            cell.text = ""
            run = cell.paragraphs[0].add_run(name)
            run.bold = True
        for row in table.rows:
            cells = grid.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = ""
                self.text_with_citations(cell.paragraphs[0], str(value))
        self.para()

    def section(self, section: Section, level: int) -> None:
        self.doc.add_heading(section.title, level=min(level, 9))
        if section.basis:
            self.para(f"Basis: {section.basis}", italic=True)
        if section.intro:
            self.para(section.intro)
        if section.narrative and section.narrative.sentences:
            self.para(" ".join(section.narrative.sentences))
        for table in section.tables:
            self.table(table)
        for item in section.items:
            self.unresolved(item, level + 1)
        for note in section.notes:
            self.para(f"Note: {note}", italic=True)
        for child in section.children:
            self.section(child, level + 1)

    def unresolved(self, item: dict[str, Any], level: int) -> None:
        title = f"[{item['finding']}] {item['title']}" if item["finding"] else item["title"]
        heading = self.doc.add_heading("", level=min(level, 9))
        self.text_with_citations(heading, title)
        flow = item["flow"]
        if flow.get("sources"):
            self.para("Sources: " + "; ".join(flow["sources"]), style="List Bullet")
        if flow.get("sink"):
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run("Sink: ")
            self.link(p, flow["sink"], url=flow.get("sink_url"))
        if flow.get("destination"):
            self.para(f"Destination: {flow['destination']}", style="List Bullet")
        for step in item["evidence"]:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run("Unresolved step: ")
            self.link(p, f"{step.get('file')}:{step.get('line')}", url=step.get("url"))
            p.add_run(f" {step.get('text', '')} ({step.get('note', '')})")
        for d in item["decisions"]:
            self.para(
                f"Low-confidence decision: {d['question']} = {d['answer']} (p = "
                f"{d['probability']:.2f})",
                style="List Bullet",
            )
        self.para("Reviewer must check:").runs[0].bold = True
        for text in item["instructions"]:
            self.para(text, style="List Number")

    def finding(self, f: dict[str, Any]) -> None:
        heading = self.doc.add_heading(f"{f['id']}: {f['title']}", level=3)
        self.bookmark(heading, bookmark_name(f["id"]))
        v = f.get("verification") or {}
        rows = [
            ["Severity", f"{f['severity']} ({', '.join(f['severity_modifiers']) or 'base'})"],
            ["Status", f["status"]],
            ["Likelihood", f["likelihood"]],
            ["Data categories", ", ".join(f["data_categories"]) or "-"],
            ["Destination class", f["destination_class"] or "-"],
            ["Dynamic verification", v.get("status", "not run")],
            ["Statutes", ", ".join(f["statute_refs"]) or "-"],
            ["Measure", f["measure"]],
        ]
        self.table(Table(["Field", "Value"], rows))
        self.para("Nodes").runs[0].bold = True
        for n in f["nodes"]:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(f"{n['id']} {n['kind']} {n['symbol']} at ")
            self.link(p, f"{n['file']}:{n['line_start']}", url=n["url"])
        if f["edges"]:
            self.para("Edges").runs[0].bold = True
            for e in f["edges"]:
                self.para(f"{e['id']} {e['from']} → {e['to']}", style="List Bullet")
                for s in e["steps"]:
                    p = self.doc.add_paragraph(style="List Bullet 2")
                    self.link(p, f"{s['file']}:{s['line']}", url=s.get("url"))
                    p.add_run(f" {s['kind']}" + (" (unresolved)" if s.get("unresolved") else ""))
        if f["decisions"]:
            rows = [
                [
                    d["id"],
                    d["target_id"],
                    d["question"],
                    d["answer"],
                    f"{d['probability']:.2f}",
                    d["provider"],
                ]
                for d in f["decisions"]
            ]
            self.table(Table(["Decision", "Target", "Question", "Answer", "p", "Provider"], rows))


def render_docx(report: Report) -> bytes:
    w = _Writer()
    run = report.run
    s = report.summary
    w.doc.add_heading(f"Data protection impact assessment: {run.repo}", level=0)
    meta = w.para(
        f"Commit {run.commit} · run {run.run_id or '-'} · generated {run.generated_at} · decision "
        f"provider {s['provider'] or '-'} · question set {s['question_set_version'] or '-'} · "
        f"threshold {s['threshold']} · narrative: {', '.join(s.get('narrative', [])) or 'template'}"
    )
    for r in meta.runs:
        r.font.color.rgb = RGBColor(0x5B, 0x61, 0x66)
    w.para(DISCLAIMER, italic=True)
    counts = ", ".join(f"{v} {k}" for k, v in sorted(s["by_severity"].items()))
    w.para(f"{s['findings']} findings: {counts}; {s['unresolved']} unresolved.")
    for section in report.sections:
        w.section(section, 1)
    w.doc.add_heading("Appendix A. Findings", level=1)
    for f in report.findings:
        w.finding(f)
    w.doc.add_heading("Appendix B. Statute references", level=1)
    for ref in report.statutes.values():
        verify = " (re-check before relying on it)" if ref.get("verify") else ""
        note = f" {ref['note']}" if ref.get("note") else ""
        w.para(
            f"{ref['id']}: {ref['citation']}, {ref['title']}.{note}{verify}", style="List Bullet"
        )
    buffer = io.BytesIO()
    w.doc.save(buffer)
    return buffer.getvalue()
